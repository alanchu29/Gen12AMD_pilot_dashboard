"""Convert the NPI dashboard workbook into the app's seed data.

Usage:
    python tools/excel_to_seed.py "<path to .xlsx>" [--out data]

Outputs:
    data/seed.json   - canonical dataset (also the Apps Script import payload)
    data/seed.js     - same data as `window.NPI_SEED = {...}` so index.html works from file://
    tools/_fixture.json - Excel's exact holiday ranges + cached dates, used by verify_engine.js

Each schedule sheet (row 7 header "# / TASK / LEAD / ... ") becomes one PFAM.
START/END formulas are translated into scheduling rules:
    E = WORKDAY(XLOOKUP(pred, A:A, F:F), lag, holidays)   -> startMode "dep"
    F = WORKDAY(E, I + endAdj, holidays)                  -> endMode "dur"
Literal dates become "manual"; anything else keeps Excel's cached value as manual.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

# Holiday sheet layout: (site, name column, date column)
HOLIDAY_COLS = [("WYMX", "A", "B"), ("WYLZ", "E", "F"), ("WYHQ", "I", "J"), ("WCZ", "M", "N"), ("WYMY", "Q", "R")]
DATE_COL_TO_SITE = {d: s for s, _, d in HOLIDAY_COLS}

RE_START = re.compile(
    r'^=WORKDAY\(\s*(?:_xlfn\.)?XLOOKUP\(\s*J(\d+)\s*,\s*A:A\s*,\s*F:F\s*,\s*"Check"\s*,\s*0\s*,\s*1\s*\)\s*,'
    r'\s*(-?\d+)\s*,\s*(.+)\)$'
)
RE_END = re.compile(r"^=WORKDAY\(\s*E(\d+)\s*,\s*I(\d+)\s*([+-]\s*\d+)?\s*,\s*(.+)\)$")
RE_HOL = re.compile(r"^(?:\[\d+\])?Holiday!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$")


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").strip()
    return "" if s.startswith("#") and s.endswith("!") else s


def parse_cal(arg, pfam_col):
    """Return (calendar rule, exact excel range) for a WORKDAY holidays argument."""
    arg = arg.strip()
    if arg == "0":
        return "none", None
    m = RE_HOL.match(arg)
    if not m:
        return "site", None
    col, r1, _, r2 = m.groups()
    rng = {"col": col, "r1": int(r1), "r2": int(r2)}
    return ("site" if col == pfam_col else DATE_COL_TO_SITE.get(col, "site")), rng


def read_holidays(ws):
    cals, cols = {}, {}
    for site, ncol, dcol in HOLIDAY_COLS:
        items, last_name, by_row = [], "", {}
        for r in range(2, ws.max_row + 1):
            name = clean(ws[f"{ncol}{r}"].value) or last_name
            last_name = name
            d = iso(ws[f"{dcol}{r}"].value)
            if d:
                items.append({"date": d, "name": name})
                by_row[r] = d
        dedup = {i["date"]: i for i in items}
        cals[site] = {"site": site, "holidays": sorted(dedup.values(), key=lambda i: i["date"])}
        cols[dcol] = by_row
    return cals, cols


def majority_col(wf):
    counts = {}
    for r in range(8, wf.max_row + 1):
        for c in "EF":
            f = wf[f"{c}{r}"].value
            if isinstance(f, str):
                for col in re.findall(r"Holiday!\$?([A-Z]+)\$?\d+", f):
                    counts[col] = counts.get(col, 0) + 1
    return max(counts, key=counts.get) if counts else None


def convert_sheet(wv, wf, idx, fixture):
    pfam_col = majority_col(wf)
    site_text = clean(wv["G1"].value)
    calendar = DATE_COL_TO_SITE.get(pfam_col) or ("WYMX" if "MX" in site_text.upper() else "WYLZ")

    tasks, wbs_to_id, pending_preds, notes_blocks = [], {}, [], []
    fx_tasks = {}
    n = 0
    for r in range(8, wv.max_row + 1):
        a, b, c = wv[f"A{r}"].value, clean(wv[f"B{r}"].value), clean(wv[f"C{r}"].value)
        e_v, f_v = wv[f"E{r}"].value, wv[f"F{r}"].value
        e_f, f_f = wf[f"E{r}"].value, wf[f"F{r}"].value
        work = wv[f"I{r}"].value
        a_s = clean(a)

        # Free-text blocks (e.g. merged "Assumption & Risk assessment")
        if a_s and not re.match(r"^\d+(\.\d+)*$", a_s) and not b:
            notes_blocks.append(a_s)
            continue
        if not b:
            continue

        is_section = (a_s and "." not in a_s) or (not a_s and work is None and not c)
        n += 1
        tid = f"t{n}"
        if is_section:
            task = {"id": tid, "type": "section", "name": b}
            if not a_s:
                task["unnumbered"] = True
            if isinstance(e_v, str) and clean(e_v):
                task["notes"] = clean(e_v)
            tasks.append(task)
            continue

        task = {
            "id": tid, "type": "task", "name": b, "lead": c,
            "startMode": "none", "pred": None, "lag": 1, "start": None,
            "workdays": work if isinstance(work, (int, float)) else None,
            "endMode": "dur", "endAdj": -1, "end": None,
            "startCal": "site", "endCal": "site",
            "pct": 0, "code": clean(wv[f"K{r}"].value), "notes": clean(wv[f"L{r}"].value),
        }
        pct = wv[f"H{r}"].value
        if isinstance(pct, (int, float)):
            # "% DONE" is percent-formatted: 0 < v <= 1 is a fraction (1 = 100%)
            task["pct"] = round(pct * 100) if 0 < pct <= 1 else int(pct)
        if a_s:
            wbs_to_id[a_s] = tid
        fx = {"row": r, "sheetWbs": a_s, "xlStart": iso(e_v), "xlEnd": iso(f_v)}

        # START
        m = RE_START.match(e_f) if isinstance(e_f, str) else None
        if m and int(m.group(1)) == r:
            task["startMode"] = "dep"
            task["lag"] = int(m.group(2))
            task["startCal"], fx["startRange"] = parse_cal(m.group(3), pfam_col)
            pending_preds.append((task, clean(wv[f"J{r}"].value)))
        elif iso(e_v):
            task["startMode"] = "manual"
            task["start"] = iso(e_v)
            if isinstance(e_f, str) and e_f.startswith("="):
                task["notes"] = (task["notes"] + "\n" if task["notes"] else "") + f"[Excel START formula {e_f}]"
        elif isinstance(e_v, str) and clean(e_v):
            task["startText"] = clean(e_v)

        # END
        m = RE_END.match(f_f) if isinstance(f_f, str) else None
        if m and int(m.group(1)) == r and int(m.group(2)) == r:
            task["endAdj"] = int(m.group(3).replace(" ", "")) if m.group(3) else 0
            task["endCal"], fx["endRange"] = parse_cal(m.group(4), pfam_col)
        elif iso(f_v):
            task["endMode"] = "manual"
            task["end"] = iso(f_v)

        if fx["xlStart"] or fx["xlEnd"]:
            task["base"] = {"start": fx["xlStart"], "end": fx["xlEnd"]}
        fx_tasks[tid] = fx
        tasks.append(task)

    for task, pred_wbs in pending_preds:
        pid = wbs_to_id.get(pred_wbs)
        task["pred"] = pid
        if not pid:
            task["predMissing"] = pred_wbs

    pfam = {
        "id": f"p{idx:03d}",
        "sheet": wv.title.strip(),
        # Exact Excel name = identity for re-imports ("x (3.5)" and "x (3.5) " are different sheets);
        # "sheet" is the display name and may be renamed in the web app.
        "xlSheet": wv.title,
        "title": clean(wv["A1"].value),
        "site": site_text.lstrip("@") or calendar,
        "calendar": calendar,
        "hidden": wv.sheet_state != "visible",
        "order": idx,
        "meta": {
            "l11pn": clean(wv["B2"].value),
            "l10pn": clean(wv["B3"].value),
            "mdm": clean(wv["C2"].value),
            "lead": clean(wv["C5"].value),
        },
        "assumptions": "\n\n".join(notes_blocks),
        "tasks": tasks,
    }
    fixture[pfam["id"]] = {"calendarCol": pfam_col, "tasks": fx_tasks}
    return pfam


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "data"))
    args = ap.parse_args()

    wbv = openpyxl.load_workbook(args.xlsx, data_only=True)
    wbf = openpyxl.load_workbook(args.xlsx)
    calendars, holiday_cols = read_holidays(wbv["Holiday"])

    pfams, fixture = [], {}
    for ws in wbf.worksheets:
        if ws["B7"].value != "TASK":
            continue
        pfams.append(convert_sheet(wbv[ws.title], ws, len(pfams) + 1, fixture))

    now = dt.datetime.now().isoformat(timespec="seconds")
    seed = {
        "schema": 1,
        "source": {"file": os.path.basename(args.xlsx), "importedAt": now},
        "calendars": calendars,
        "pfams": pfams,
    }
    os.makedirs(args.out, exist_ok=True)
    with open(os.path.join(args.out, "seed.json"), "w", encoding="utf-8") as f:
        json.dump(seed, f, ensure_ascii=False, indent=1)
    with open(os.path.join(args.out, "seed.js"), "w", encoding="utf-8") as f:
        f.write("// Generated by tools/excel_to_seed.py - do not edit by hand.\n")
        f.write("window.NPI_SEED = " + json.dumps(seed, ensure_ascii=False) + ";\n")
    with open(os.path.join(os.path.dirname(__file__), "_fixture.json"), "w", encoding="utf-8") as f:
        json.dump({"holidayCols": holiday_cols, "pfams": fixture}, f)

    visible = sum(1 for p in pfams if not p["hidden"])
    ntasks = sum(1 for p in pfams for t in p["tasks"] if t["type"] == "task")
    missing = [(p["sheet"], t["name"], t["predMissing"]) for p in pfams for t in p["tasks"] if t.get("predMissing")]
    print(f"PFAMs: {len(pfams)} ({visible} visible), tasks: {ntasks}")
    for site, cal in calendars.items():
        print(f"  calendar {site}: {len(cal['holidays'])} holidays")
    if missing:
        print(f"Unresolved predecessors: {len(missing)}")
        for m in missing[:20]:
            print("   ", m)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
